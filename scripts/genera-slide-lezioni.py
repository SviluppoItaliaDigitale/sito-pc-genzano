#!/usr/bin/env python3
# Genera le slide per chi conduce le lezioni nelle scuole (PowerPoint con le note per chi
# parla, più PDF) da data/slide_lezioni.yaml, con la grafica degli Appunti illustrati.
"""
Uso:
  python3 scripts/genera-slide-lezioni.py              # scrive i .pptx
  python3 scripts/genera-slide-lezioni.py --pdf        # scrive i .pptx e i .pdf
  python3 scripts/genera-slide-lezioni.py --verifica   # controlla i file già scritti
  python3 scripts/genera-slide-lezioni.py --solo 2     # solo la lezione 2

Uscita: static/formazione/slide-lezioni/percorso-<p>-lezione-<n>-<slug>.pptx (e .pdf).
Dopo la scrittura allinea da solo le dimensioni dei file citate nelle pagine
content/formazione/percorsi-didattici/_index.md e content/area-download/_index.md.

Dipendenze: python-pptx, pyyaml, pillow, cairosvg, fonttools, brotli (per leggere il
carattere WOFF2 degli Appunti illustrati); per --pdf anche LibreOffice e pymupdf.

Scelte che non vanno cambiate senza motivo:
- I titoli sono scritti con il carattere a mano degli Appunti illustrati
  (Edu NSW ACT Cursive, licenza SIL OFL 1.1), incorporato nel PowerPoint nel formato
  che PowerPoint usa per i caratteri incorporati (Embedded OpenType); il testo è Arial.
  Il PDF incorpora il carattere da sé: LibreOffice lo trova in una cartella temporanea.
- Testo delle slide di almeno 24 punti, titoli di almeno 36; ogni immagine ha il suo
  testo alternativo o è segnata come decorativa; ogni slide ha il titolo nel segnaposto
  del titolo (anche per chi usa un lettore di schermo) e le note per chi parla.
- Esattamente quattro loghi per lezione: il Gruppo solo in copertina, le tre
  affiliazioni solo nell'ultima slide, con il codice E10435833 accanto al Quality Label.
- Nessuna proprietà del documento cita strumenti: autore e firma sono del Gruppo.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import cairosvg
import yaml
from fontTools.ttLib import TTFont
from fontTools.varLib import instancer
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DATI = ROOT / "data" / "slide_lezioni.yaml"
USCITA = ROOT / "static" / "formazione" / "slide-lezioni"
PITTOGRAMMI = ROOT / "static" / "pittogrammi"
ILLUSTRAZIONI = ROOT / "static" / "formazione" / "illustrazioni-udl"
CARATTERE_WOFF2 = ROOT / "static" / "formazione" / "schede-stampabili" / "assets" / "edu-cursive-400.woff2"
LOGO_GRUPPO = ROOT / "static" / "images" / "logo-pc-genzano-hires.png"
AFFILIAZIONI = [
    (ROOT / "static" / "images" / "quality-label-esc.png", "Logo Quality Label del Corpo europeo di solidarietà"),
    (ROOT / "static" / "images" / "logo-fepivol.png", "Logo del Coordinamento FE.PI.VOL., Federazione Pronto Intervento Volontariato"),
    (ROOT / "static" / "images" / "logo-snpc-volontariato.png", "Logo del Servizio nazionale della protezione civile, Volontariato"),
]
CODICE_ESC = "E10435833"
PAGINE_CON_LINK = [
    ROOT / "content" / "formazione" / "percorsi-didattici" / "_index.md",
    ROOT / "content" / "area-download" / "_index.md",
]

SCRITTURA = "Edu NSW ACT Cursive"
TESTO = "Arial"
MISURA_ARIAL = Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf")
MISURA_ARIAL_B = Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf")

# Colori degli Appunti illustrati (assets/appunti.css e scheda-print.css).
INK = "003366"        # inchiostro blu: titoli, bordi
NERO = "212529"       # testo
CARTA = "FFFDF7"      # foglio
GIALLO = "FFF0A8"     # foglietto giallo, evidenziatore «azione»
AZZURRO = "DCEEFA"    # foglietto azzurro, evidenziatore «luogo»
ROSA = "F8DFE5"       # riquadro errori, evidenziatore «attenzione»
VERDE_CHIARO = "EDF7F0"
VERDE = "146C43"
ROSSO = "B02A37"
ORO = "B45309"        # frecce
PENNELLO = "FFE699"   # pennellata sotto i titoli
BIANCO = "FFFFFF"
GRIGIO_BORDO = "ADB5BD"
EVIDENZA = {"azione": GIALLO, "luogo": AZZURRO, "attenzione": ROSA}
TINTE = [AZZURRO, GIALLO, ROSA, VERDE_CHIARO]

LARGHEZZA, ALTEZZA = 13.333, 7.5
MARGINE = 0.7
MIN_TESTO, MIN_TITOLO = 24, 36

# Disegni degli Appunti illustrati (scripts/genera-appunti-illustrati.py): pennellata, freccia, icone.
SVG_PENNELLO = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 16"><path d="M4 10Q100 0 295 7L287 13Q120 7 9 15Z" fill="#{c}"/></svg>'
SVG_FRECCIA = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 48 20"><path d="M2 12Q22 5 44 10m-9-7 9 7-10 7" '
               'fill="none" stroke="#{c}" stroke-width="2.5" stroke-linecap="round" stroke-linejoin="round"/></svg>')
ICONE = {
    "avviso": '<path d="m14 2 13 24H1zM14 10v8m0 3v1"/>',
    "lampada": '<path d="M9 20c0-4-5-5-5-10a10 10 0 0 1 20 0c0 5-5 6-5 10zM9 24h10M12 27h4"/>',
}
SVG_ICONA = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 28 28" fill="none" stroke="#{c}" stroke-width="1.8" '
             'stroke-linecap="round" stroke-linejoin="round">{p}</svg>')

DECORATIVO_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"
DECORATIVO_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
PAROLE_VIETATE = ("python-pptx", "steve canny", "claude", "chatgpt", "openai", "anthropic", "codex", "gemini")


class ErroreSlide(Exception):
    pass


# ───────────────────────────── risorse ─────────────────────────────

class Risorse:
    """Immagini e carattere preparati una volta, in una cartella temporanea."""

    def __init__(self, cartella: Path):
        self.dir = cartella
        self.cache: dict[str, Path] = {}
        self.ttf = self._carattere()
        self.misura_scrittura = Misura(self.ttf)
        self.misura_testo = Misura(MISURA_ARIAL)
        self.misura_grassetto = Misura(MISURA_ARIAL_B)

    def _carattere(self) -> Path:
        if not CARATTERE_WOFF2.exists():
            raise SystemExit(f"Manca il carattere degli Appunti illustrati: {CARATTERE_WOFF2.relative_to(ROOT)}")
        font = TTFont(CARATTERE_WOFF2)
        font.flavor = None
        if "fvar" in font:
            font = instancer.instantiateVariableFont(font, {"wght": 400}, updateFontNames=True)
        out = self.dir / "EduNSWACTCursive-Regular.ttf"
        font.save(out)
        return out

    def svg(self, nome: str, codice: str, larghezza: int) -> Path:
        chiave = f"{nome}-{larghezza}"
        if chiave not in self.cache:
            out = self.dir / f"{chiave}.png"
            cairosvg.svg2png(bytestring=codice.encode("utf-8"), write_to=str(out), output_width=larghezza)
            self.cache[chiave] = out
        return self.cache[chiave]

    def pennello(self) -> Path:
        return self.svg("pennello", SVG_PENNELLO.format(c=PENNELLO), 1200)

    def freccia(self) -> Path:
        return self.svg("freccia", SVG_FRECCIA.format(c=ORO), 480)

    def icona(self, nome: str, colore: str) -> Path:
        return self.svg(f"icona-{nome}-{colore}", SVG_ICONA.format(c=colore, p=ICONE[nome]), 280)

    def illustrazione(self, nome: str, ritaglio: str | None) -> tuple[Path, float]:
        sorgente = ILLUSTRAZIONI / nome
        if not sorgente.exists():
            raise ErroreSlide(f"Illustrazione inesistente: {sorgente.relative_to(ROOT)}")
        testo = sorgente.read_text(encoding="utf-8")
        if ritaglio:
            x, y, w, h = (float(v) for v in ritaglio.split())
            testo = re.sub(r'viewBox="[^"]*"', f'viewBox="{ritaglio}"', testo, count=1)
            testo = re.sub(r'(<svg[^>]*?)\swidth="[^"]*"', rf'\1 width="{w:g}"', testo, count=1)
            testo = re.sub(r'(<svg[^>]*?)\sheight="[^"]*"', rf'\1 height="{h:g}"', testo, count=1)
        else:
            m = re.search(r'viewBox="([^"]*)"', testo)
            x, y, w, h = (float(v) for v in m.group(1).split())
        out = self.dir / f"udl-{sorgente.stem}.png"
        cairosvg.svg2png(bytestring=testo.encode("utf-8"), write_to=str(out), output_width=1800)
        return out, w / h

    @staticmethod
    def pittogramma(rel: str) -> Path:
        p = PITTOGRAMMI / rel
        if not p.exists():
            raise ErroreSlide(f"Pittogramma inesistente: {p.relative_to(ROOT)}")
        return p


class Misura:
    """Larghezze dei caratteri lette dal file, per sapere se una riga ci sta."""

    def __init__(self, percorso: Path):
        font = TTFont(percorso)
        self.upem = font["head"].unitsPerEm
        self.cmap = font.getBestCmap()
        self.hmtx = font["hmtx"]
        self.ripiego = self.cmap.get(ord("n"))

    def larghezza(self, testo: str, punti: float) -> float:
        totale = 0
        for ch in testo:
            g = self.cmap.get(ord(ch), self.ripiego)
            totale += self.hmtx[g][0]
        return totale / self.upem * punti / 72


# ───────────────────────────── testo ─────────────────────────────

def tipografia(s: str) -> str:
    """Apostrofo tipografico, come negli Appunti illustrati."""
    return s.replace("'", "’")


def pezzi(riga) -> list[tuple[str, str | None]]:
    """Riga «ricca» → [(testo, evidenza)], come in genera-appunti-illustrati.py."""
    if isinstance(riga, str):
        return [(tipografia(riga), None)]
    out = []
    for parte in riga:
        if isinstance(parte, str):
            out.append((tipografia(parte), None))
        else:
            if parte.get("evidenza") not in EVIDENZA:
                raise ErroreSlide(f"Evidenziatore sconosciuto: {parte.get('evidenza')!r}")
            out.append((tipografia(parte["testo"]), parte["evidenza"]))
    return out


def testo_piano(riga) -> str:
    return "".join(t for t, _ in pezzi(riga))


def righe_a_capo(testo: str, larghezza: float, punti: float, misura: Misura) -> list[str]:
    parole = testo.split()
    righe, corrente = [], ""
    for parola in parole:
        prova = parola if not corrente else corrente + " " + parola
        if misura.larghezza(prova, punti) <= larghezza or not corrente:
            corrente = prova
        else:
            righe.append(corrente)
            corrente = parola
    if corrente:
        righe.append(corrente)
    return righe or [""]


class Risorsa:
    pass


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


def segna_decorativo(elemento_cnvpr):
    ext_lst = elemento_cnvpr.find(qn("a:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(elemento_cnvpr, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", DECORATIVO_URI)
    dec = etree.SubElement(ext, f"{{{DECORATIVO_NS}}}decorative", nsmap={"adec": DECORATIVO_NS})
    dec.set("val", "1")


def cnvpr(shape):
    el = shape._element
    for tag in ("p:nvSpPr", "p:nvPicPr", "p:nvCxnSpPr"):
        nv = el.find(qn(tag))
        if nv is not None:
            return nv.find(qn("p:cNvPr"))
    raise ErroreSlide("Forma senza proprietà non visive")


def imposta_run(run, testo, punti, colore, carattere=TESTO, grassetto=False, evidenza=None, link=None):
    run.text = testo
    f = run.font
    f.size = Pt(punti)
    f.bold = grassetto
    f.name = carattere
    f.color.rgb = rgb(colore)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "it-IT")
    if evidenza:
        hl = etree.Element(qn("a:highlight"))
        etree.SubElement(hl, qn("a:srgbClr")).set("val", EVIDENZA[evidenza])
        rpr.find(qn("a:latin")).addprevious(hl)
    if link:
        run.hyperlink.address = link
        f.underline = True


def puntato(paragrafo, colore=INK, rientro=0.42):
    ppr = paragrafo._p.get_or_add_pPr()
    ppr.set("marL", str(Inches(rientro)))
    ppr.set("indent", str(-Inches(rientro)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for vecchio in ppr.findall(qn(tag)):
            ppr.remove(vecchio)
    clr = etree.SubElement(ppr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", colore)
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", TESTO)
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


def interlinea(paragrafo, percento: int):
    ppr = paragrafo._p.get_or_add_pPr()
    for vecchio in ppr.findall(qn("a:lnSpc")):
        ppr.remove(vecchio)
    ln = etree.Element(qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPct")).set("val", str(percento * 1000))
    ppr.insert(0, ln)


# ───────────────────────────── costruttore ─────────────────────────────

class Lezione:
    LH_TESTO = 1.2      # altezza di riga di Arial, in multipli del corpo
    LH_SCRITTURA = 1.95  # altezza di riga del carattere a mano (metriche Windows)

    def __init__(self, dati: dict, lezione: dict, risorse: Risorse):
        self.dati = dati
        self.lez = lezione
        self.r = risorse
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(LARGHEZZA), Inches(ALTEZZA)
        self._prepara_modello()
        self.layout = self.prs.slide_layouts[5]  # «Solo titolo»
        self.indice = 0

    # ── modello: sfondo, tema, segnaposto 16:9 ──
    def _prepara_modello(self):
        master = self.prs.slide_master
        master.background.fill.solid()
        master.background.fill.fore_color.rgb = rgb(CARTA)
        fattore = LARGHEZZA / 10.0
        for contenitore in [master, *self.prs.slide_layouts]:
            for sh in contenitore.shapes:
                if sh.left is not None and sh.width is not None:
                    sh.left = Emu(int(sh.left * fattore))
                    sh.width = Emu(int(sh.width * fattore))
        tema = master.part.part_related_by("http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
        radice = etree.fromstring(tema.blob)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        radice.find(".//a:fontScheme", ns).set("name", "Appunti illustrati")
        radice.find(".//a:majorFont/a:latin", ns).set("typeface", SCRITTURA)
        radice.find(".//a:minorFont/a:latin", ns).set("typeface", TESTO)
        schema = radice.find(".//a:clrScheme", ns)
        schema.set("name", "Appunti illustrati")
        colori = {"dk1": NERO, "lt1": BIANCO, "dk2": INK, "lt2": CARTA, "accent1": INK, "accent2": ORO,
                  "accent3": VERDE, "accent4": ROSSO, "accent5": "0369A1", "accent6": "6F42C1",
                  "hlink": INK, "folHlink": "4A2C7F"}
        for nome, valore in colori.items():
            el = schema.find(f"a:{nome}", ns)
            for figlio in list(el):
                el.remove(figlio)
            etree.SubElement(el, qn("a:srgbClr")).set("val", valore)
        tema._blob = etree.tostring(radice, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ── forme di base ──
    def nuova_slide(self):
        s = self.prs.slides.add_slide(self.layout)
        self.indice += 1
        return s

    def _senza_stile(self, shape):
        stile = shape._element.find(qn("p:style"))
        if stile is not None:
            shape._element.remove(stile)

    def rettangolo(self, slide, x, y, w, h, fondo=None, bordo=None, spessore=1.5, tratteggio=False,
                   raggio=0.0, forma=None, rotazione=0.0, decorativo=True, nome="Riquadro", ombra=False):
        tipo = forma or (MSO_SHAPE.ROUNDED_RECTANGLE if raggio else MSO_SHAPE.RECTANGLE)
        sh = slide.shapes.add_shape(tipo, Inches(x), Inches(y), Inches(w), Inches(h))
        self._senza_stile(sh)
        sh.name = nome
        if raggio and tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = min(0.5, raggio / min(w, h))
        if fondo:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fondo)
        else:
            sh.fill.background()
        if bordo:
            sh.line.color.rgb = rgb(bordo)
            sh.line.width = Pt(spessore)
            if tratteggio:
                sh.line.dash_style = MSO_LINE.DASH
        else:
            sh.line.fill.background()
        if rotazione:
            sh.rotation = rotazione
        if ombra:
            sppr = sh._element.spPr
            eff = etree.SubElement(sppr, qn("a:effectLst"))
            o = etree.SubElement(eff, qn("a:outerShdw"))
            o.set("blurRad", str(Emu(Pt(5)))); o.set("dist", str(Emu(Pt(2)))); o.set("dir", "5400000")
            o.set("algn", "t"); o.set("rotWithShape", "0")
            c = etree.SubElement(o, qn("a:srgbClr")); c.set("val", "000000")
            etree.SubElement(c, qn("a:alpha")).set("val", "22000")
        if decorativo:
            segna_decorativo(cnvpr(sh))
        return sh

    def immagine(self, slide, percorso: Path, x, y, w=None, h=None, alt=None, decorativo=False,
                 rotazione=0.0, nome="Immagine", contieni=True):
        """Inserisce l'immagine; con «contieni» la adatta al riquadro w×h senza deformarla."""
        if contieni and w and h:
            with Image.open(percorso) as im:
                iw, ih = im.size
            scala = min(w / iw, h / ih)
            nw, nh = iw * scala, ih * scala
            x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
        pic = slide.shapes.add_picture(str(percorso), Inches(x), Inches(y),
                                       Inches(w) if w else None, Inches(h) if h else None)
        pic.name = nome
        el = cnvpr(pic)
        if decorativo:
            el.set("descr", "")
            segna_decorativo(el)
        else:
            if not alt:
                raise ErroreSlide(f"Immagine senza testo alternativo: {percorso.name}")
            el.set("descr", tipografia(alt))
        if rotazione:
            pic.rotation = rotazione
        return pic

    def casella(self, slide, x, y, w, h, nome="Testo", ancoraggio=MSO_ANCHOR.TOP, margini=(0.08, 0.05)):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.name = nome
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = ancoraggio
        tf.margin_left = tf.margin_right = Inches(margini[0])
        tf.margin_top = tf.margin_bottom = Inches(margini[1])
        return tb

    def forma_con_testo(self, shape, ancoraggio=MSO_ANCHOR.MIDDLE, margini=(0.15, 0.08)):
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = ancoraggio
        tf.margin_left = tf.margin_right = Inches(margini[0])
        tf.margin_top = tf.margin_bottom = Inches(margini[1])
        # una forma con testo non è decorativa: tolgo il segno se c'era
        el = cnvpr(shape)
        ext = el.find(qn("a:extLst"))
        if ext is not None:
            el.remove(ext)
        return tf

    # ── misure ──
    def altezza_paragrafi(self, righe, larghezza, punti, grassetto=False, spazio_dopo=10, rientro=0.0):
        misura = self.r.misura_grassetto if grassetto else self.r.misura_testo
        tot = 0.0
        for riga in righe:
            n = len(righe_a_capo(testo_piano(riga), larghezza - rientro, punti, misura))
            tot += n * punti * self.LH_TESTO / 72
        tot += (len(righe) - 1) * spazio_dopo / 72
        return tot

    def corpo_adatto(self, righe, larghezza, altezza, massimo=28, grassetto=False, spazio_dopo=10, rientro=0.0):
        misura = self.r.misura_grassetto if grassetto else self.r.misura_testo
        parole = [p for r in righe for p in testo_piano(r).split()]
        for punti in range(massimo, MIN_TESTO - 1, -1):
            # nessuna parola deve essere più larga della riga: verrebbe spezzata a metà
            if any(misura.larghezza(p, punti) > larghezza - rientro for p in parole):
                continue
            if self.altezza_paragrafi(righe, larghezza, punti, grassetto, spazio_dopo, rientro) <= altezza:
                return punti
        testo = " / ".join(testo_piano(r) for r in righe)
        raise ErroreSlide(f"Lezione {self.lez['numero']}, slide {self.indice}: il testo non entra "
                          f"nemmeno a {MIN_TESTO} punti: {testo[:90]}…")

    def scrivi(self, tf, righe, punti, colore=NERO, grassetto=False, elenco=False, spazio_dopo=10,
               allineamento=PP_ALIGN.LEFT):
        tf.clear()
        for i, riga in enumerate(righe):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = allineamento
            p.space_after = Pt(spazio_dopo) if i < len(righe) - 1 else Pt(0)
            if elenco:
                puntato(p)
            for testo, ev in pezzi(riga):
                imposta_run(p.add_run(), testo, punti, colore, grassetto=grassetto, evidenza=ev)

    # ── titolo con pennellata ──
    def titolo(self, slide, testo, x=MARGINE, y=0.25, w=12.0, punti=46, righe_max=1, interlinea_pct=100):
        testo = tipografia(testo)
        misura = self.r.misura_scrittura
        while True:
            righe = righe_a_capo(testo, w - 0.2, punti, misura)
            if len(righe) <= righe_max or punti <= MIN_TITOLO:
                break
            punti -= 2
        if len(righe) > righe_max:
            raise ErroreSlide(f"Lezione {self.lez['numero']}, slide {self.indice}: titolo troppo lungo: {testo}")
        riga_h = punti * self.LH_SCRITTURA * interlinea_pct / 100 / 72
        h = riga_h * len(righe) + 0.1
        t = slide.shapes.title
        t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(h)
        tf = t.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        tf.margin_left = tf.margin_right = Inches(0.1)
        tf.margin_top = tf.margin_bottom = Inches(0.05)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        if interlinea_pct != 100:
            interlinea(p, interlinea_pct)
        for i, riga in enumerate(righe):
            if i:
                p.add_line_break()
            imposta_run(p.add_run(), riga, punti, INK, carattere=SCRITTURA)
        # pennellata sotto l'ultima riga, dietro al testo (come negli Appunti illustrati)
        base = y + h - 0.05 - punti * 0.47 / 72
        larga = min(misura.larghezza(righe[-1], punti) + 0.25, w - 0.1)
        pen = self.immagine(slide, self.r.pennello(), x + 0.02, base + 0.01, larga, 0.22,
                            decorativo=True, nome="Pennellata", contieni=False)
        t._element.addprevious(pen._element)
        return y + h

    # ── elementi ricorrenti ──
    def cartoncino(self, slide, file, alt, x, y, lato, rotazione=0.0, fondo=BIANCO, nome="Pittogramma"):
        self.rettangolo(slide, x, y, lato, lato, fondo=fondo, bordo=GRIGIO_BORDO, spessore=0.75,
                        rotazione=rotazione, nome="Cartoncino", ombra=True)
        pad = lato * 0.07
        return self.immagine(slide, self.r.pittogramma(file), x + pad, y + pad, lato - 2 * pad, lato - 2 * pad,
                             alt=alt, rotazione=rotazione, nome=nome)

    def foglietto(self, slide, x, y, w, h, colore, rotazione, testo, punti=32, carattere=SCRITTURA,
                  colore_testo=INK, nome="Foglietto"):
        sh = self.rettangolo(slide, x, y, w, h, fondo=colore, rotazione=rotazione, nome=nome, ombra=True,
                             decorativo=False)
        tf = self.forma_con_testo(sh, margini=(0.12, 0.04))
        tf.word_wrap = False  # un foglietto porta una parola o due: non va a capo
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        imposta_run(p.add_run(), tipografia(testo), punti, colore_testo, carattere=carattere)
        return sh

    def linguetta(self, slide, x, y, testo, fondo, bordo=INK, punti=30, icona=None):
        testo = tipografia(testo)
        larg = self.r.misura_scrittura.larghezza(testo, punti) * 1.12 + 0.7 + (0.6 if icona else 0)
        h = punti * 1.55 / 72
        sh = self.rettangolo(slide, x, y, larg, h, fondo=fondo, bordo=bordo, spessore=1.25,
                             forma=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, nome="Linguetta", decorativo=False)
        sh.adjustments[0] = 0.28
        tf = self.forma_con_testo(sh, margini=(0.2 + (0.55 if icona else 0), 0.0))
        tf.word_wrap = False  # un'etichetta non va mai a capo
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        imposta_run(p.add_run(), testo, punti, INK, carattere=SCRITTURA)
        if icona:
            self.immagine(slide, icona, x + 0.15, y + (h - 0.48) / 2, 0.48, 0.48, decorativo=True, nome="Icona")
        return larg, h

    def numero_tondo(self, slide, x, y, d, numero, punti=30):
        sh = self.rettangolo(slide, x, y, d, d, fondo=BIANCO, bordo=INK, spessore=2, forma=MSO_SHAPE.OVAL,
                             nome="Numero", decorativo=False)
        tf = self.forma_con_testo(sh, margini=(0.0, 0.0))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        imposta_run(p.add_run(), str(numero), punti, INK, carattere=SCRITTURA)

    def blocco_testo(self, slide, righe, x, y, w, h, massimo=28, elenco=True, grassetto=False, colore=NERO,
                     ancoraggio=MSO_ANCHOR.TOP, spazio_dopo=12, nome="Testo"):
        rientro = 0.42 if elenco else 0.0
        punti = self.corpo_adatto(righe, w - 0.16, h - 0.1, massimo, grassetto, spazio_dopo, rientro)
        tb = self.casella(slide, x, y, w, h, nome=nome, ancoraggio=ancoraggio)
        self.scrivi(tb.text_frame, righe, punti, colore, grassetto, elenco, spazio_dopo)
        return tb

    # ── note per chi parla ──
    def note(self, slide, note: dict, extra: list[str] | None = None):
        if not note or "dire" not in note or "tempo" not in note:
            raise ErroreSlide(f"Lezione {self.lez['numero']}, slide {self.indice}: note incomplete")
        righe = [f"Cosa dire: {note['dire']}"]
        if note.get("domanda"):
            righe.append(f"Domanda alla classe: {note['domanda']}")
        minuti = int(note["tempo"])
        righe.append(f"Tempo: {minuti} minut{'o' if minuti == 1 else 'i'}" +
                     (" (slide per l'insegnante, non conta nell'ora)" if minuti == 0 else ""))
        if note.get("attenzione"):
            righe.append(f"Attenzione: {note['attenzione']}")
        righe += extra or []
        tf = slide.notes_slide.notes_text_frame
        tf.clear()
        for i, r in enumerate(righe):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = tipografia(r)
            run.font.size = Pt(14)
            run._r.get_or_add_rPr().set("lang", "it-IT")

    # ───────────────────────── tipi di slide ─────────────────────────

    def copertina(self, s):
        sl = self.nuova_slide()
        perc = self.dati["percorso"]
        tb = self.casella(sl, MARGINE, 0.45, 8.6, 1.0, nome="Percorso")
        self.scrivi(tb.text_frame, [f"Percorso {perc['numero']} · {perc['titolo']}", perc["fascia"]], 24,
                    INK, spazio_dopo=2)
        tb.text_frame.paragraphs[0].runs[0].font.bold = True
        fine_titolo = self.titolo(sl, self.lez["titolo"], x=MARGINE, y=1.45, w=7.9, punti=64, righe_max=3,
                                  interlinea_pct=78)
        y = max(fine_titolo + 0.45, 4.55)
        totale = len(self.dati["lezioni"])
        quale = f"Lezione {self.lez['numero']} di {totale}"
        largo = max(3.25, self.r.misura_scrittura.larghezza(quale, 34) * 1.12 + 0.6)
        self.foglietto(sl, MARGINE + 0.05, y, largo, 0.9, GIALLO, -3, quale, punti=34, nome="Lezione")
        self.foglietto(sl, MARGINE + 0.45 + largo, y + 0.05, 2.1, 0.85, AZZURRO, 2, "Un'ora", punti=34,
                       nome="Durata")
        self.immagine(sl, LOGO_GRUPPO, 11.45, 0.35, 1.45, 1.45,
                      alt="Logo del Gruppo Comunale Volontari di Protezione Civile di Genzano di Roma",
                      nome="Logo del Gruppo")
        disegni = self.lez["disegni"]
        if len(disegni) == 1:
            d = disegni[0]
            self.cartoncino(sl, d["file"], d["alt"], 9.05, 2.25, 3.4, rotazione=2)
        else:
            posti = [(8.75, 1.95, 2.35, -4), (10.55, 2.75, 2.35, 3), (9.1, 4.3, 2.05, -2)]
            for d, (x, y2, lato, rot) in zip(disegni, posti):
                self.cartoncino(sl, d["file"], d["alt"], x, y2, lato, rotazione=rot)
        firma = self.casella(sl, MARGINE, 6.55, 12.0, 0.55, nome="Firma")
        self.scrivi(firma.text_frame, [self.dati["firma"]], 24, INK)
        self.note(sl, s["note"])

    def obiettivi(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s.get("titolo", "Oggi scopriamo"))
        self.rettangolo(sl, MARGINE, 1.75, 8.55, 4.95, bordo=INK, tratteggio=True, raggio=0.12, nome="Riquadro")
        punti = s["punti"]
        n = len(punti)
        alto = 4.55 / n
        for i, riga in enumerate(punti):
            y = 1.95 + i * alto
            self.numero_tondo(sl, MARGINE + 0.3, y + (alto - 0.8) / 2, 0.8, i + 1, punti=34)
            self.blocco_testo(sl, [riga], MARGINE + 1.35, y, 6.95, alto, elenco=False,
                              ancoraggio=MSO_ANCHOR.MIDDLE, nome=f"Obiettivo {i + 1}")
        im = s["immagine"]
        self.cartoncino(sl, im["file"], im["alt"], 9.75, 2.2, 2.9, rotazione=2)
        self.note(sl, s["note"])

    def storia(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        im = s["immagine"]
        self.cartoncino(sl, im["file"], im["alt"], 0.85, 1.95, 2.85, rotazione=-2)
        if im.get("nome"):
            self.foglietto(sl, 1.15, 5.05, 2.3, 0.78, AZZURRO, 2, im["nome"], punti=32, nome="Nome")
        foglio = self.rettangolo(sl, 4.15, 1.8, 8.45, 3.45, fondo=GIALLO, rotazione=-1, ombra=True,
                                 nome="Foglietto della storia")
        self.blocco_testo(sl, s["testo"], 4.45, 1.95, 7.9, 3.15, elenco=False, ancoraggio=MSO_ANCHOR.MIDDLE,
                          nome="Storia")
        self.rettangolo(sl, 4.15, 5.6, 8.45, 1.3, bordo=INK, tratteggio=True, raggio=0.12, nome="Riquadro")
        self.immagine(sl, self.r.icona("lampada", INK), 4.4, 5.85, 0.8, 0.8, decorativo=True, nome="Icona")
        self.blocco_testo(sl, [s["domanda"]], 5.3, 5.65, 7.2, 1.2, elenco=False, grassetto=True,
                          colore=INK, ancoraggio=MSO_ANCHOR.MIDDLE, nome="Domanda")
        self.note(sl, s["note"])

    def testo_immagine(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        immagini = s["immagini"]
        sinistra = s.get("lato") == "sinistra"
        bx = 4.55 if sinistra else MARGINE
        ix = MARGINE + 0.15 if sinistra else 9.3
        self.rettangolo(sl, bx, 1.75, 8.1, 4.95, bordo=INK, tratteggio=True, raggio=0.12, nome="Riquadro")
        self.blocco_testo(sl, s["punti"], bx + 0.3, 1.95, 7.55, 4.55, ancoraggio=MSO_ANCHOR.MIDDLE)
        if len(immagini) == 1:
            im = immagini[0]
            self.cartoncino(sl, im["file"], im["alt"], ix, 2.25, 3.2, rotazione=-2 if sinistra else 2)
        else:
            a, b = immagini[:2]
            self.cartoncino(sl, a["file"], a["alt"], ix, 1.85, 2.45, rotazione=-2)
            self.cartoncino(sl, b["file"], b["alt"], ix + 0.6, 4.35, 2.3, rotazione=3)
        self.note(sl, s["note"])

    def illustrazione(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        png, rapporto = self.r.illustrazione(s["svg"], s.get("ritaglio"))
        righe = s["punti"]
        h_testo = self.altezza_paragrafi(righe, 11.6 - 0.58, 26, spazio_dopo=8, rientro=0.42) + 0.2
        h_max = 7.2 - 1.75 - 0.25 - h_testo
        w = min(11.0, h_max * rapporto)
        h = w / rapporto
        x = (LARGHEZZA - w) / 2
        self.rettangolo(sl, x - 0.12, 1.7, w + 0.24, h + 0.24, fondo=BIANCO, bordo=INK, tratteggio=True,
                        raggio=0.1, nome="Cornice")
        self.immagine(sl, png, x, 1.82, w, h, alt=s["alt"], nome="Illustrazione")
        y = 1.7 + h + 0.24 + 0.15
        self.blocco_testo(sl, righe, 0.85, y, 11.6, 7.25 - y, massimo=26, spazio_dopo=8)
        self.note(sl, s["note"])

    def griglia(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        voci = s["voci"]
        n = len(voci)
        gap = 0.4
        w = (12.0 - gap * (n - 1)) / n
        h = 3.85 if s.get("nota") else 4.6
        rot = [-2, 1.5, -1.5, 2]
        for i, v in enumerate(voci):
            x = MARGINE + i * (w + gap)
            self.rettangolo(sl, x, 1.95, w, h, fondo=TINTE[(i + 1) % 4], rotazione=rot[i % 4], ombra=True,
                            nome="Foglietto")
            lato = min(w - 0.6, h - 1.55, 2.6)
            self.immagine(sl, self.r.pittogramma(v["file"]), x + (w - lato) / 2, 2.15, lato, lato, alt=v["alt"],
                          rotazione=rot[i % 4], nome="Pittogramma")
            self.blocco_testo(sl, [v["testo"]], x + 0.1, 2.25 + lato, w - 0.2, h - lato - 0.4, massimo=28,
                              elenco=False, grassetto=True, colore=INK, nome=f"Voce {i + 1}")
            for p in sl.shapes[-1].text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
        if s.get("nota"):
            self.blocco_testo(sl, [s["nota"]], MARGINE, 6.05, 12.0, 1.15, massimo=26, elenco=False,
                              ancoraggio=MSO_ANCHOR.MIDDLE, nome="Nota")
        self.note(sl, s["note"])

    def sequenza(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        passi = s["passi"]
        n = len(passi)
        gap = 0.45
        w = (12.0 - gap * (n - 1)) / n
        con_immagini = any("file" in p for p in passi)
        h = 4.85 if con_immagini else 3.55
        y0 = 1.85 if con_immagini else 2.45
        for i, p in enumerate(passi):
            x = MARGINE + i * (w + gap)
            self.rettangolo(sl, x, y0, w, h, fondo=TINTE[i % 4], bordo=INK, spessore=1.25, raggio=0.25,
                            nome="Passo")
            et = self.casella(sl, x + 0.05, y0 + 0.05, w - 0.1, 0.95, nome="Etichetta", ancoraggio=MSO_ANCHOR.MIDDLE)
            par = et.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.CENTER
            imposta_run(par.add_run(), tipografia(str(p["etichetta"])), 38, INK, carattere=SCRITTURA)
            y = y0 + 1.0
            if "file" in p:
                lato = min(w - 0.7, 1.55)
                self.immagine(sl, self.r.pittogramma(p["file"]), x + (w - lato) / 2, y, lato, lato, alt=p["alt"],
                              nome="Pittogramma")
                y += lato + 0.12
            resto = y0 + h - y - 0.12
            corpo_t = self.corpo_adatto([p["titolo"]], w - 0.26, 9, massimo=26, grassetto=True)
            titolo_h = self.altezza_paragrafi([p["titolo"]], w - 0.26, corpo_t, grassetto=True) + 0.12
            tt = self.blocco_testo(sl, [p["titolo"]], x + 0.05, y, w - 0.1, titolo_h, massimo=corpo_t, elenco=False,
                                   grassetto=True, colore=INK, nome="Passo, titolo")
            tt.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            tx = self.blocco_testo(sl, [p["testo"]], x + 0.05, y + titolo_h, w - 0.1, resto - titolo_h, massimo=24,
                                   elenco=False, nome="Passo, testo")
            tx.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            if i < n - 1:
                self.immagine(sl, self.r.freccia(), x + w + 0.04, y0 + h / 2 - 0.12, gap - 0.08, 0.24,
                              decorativo=True, nome="Freccia", contieni=False)
        self.note(sl, s["note"])

    def si_no(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        lati = [(s["si"], MARGINE, 7.05, VERDE_CHIARO, VERDE, "arasaac/si.png", "Simbolo: sì"),
                (s["no"], 8.15, 4.55, ROSA, ROSSO, "arasaac/no.png", "Simbolo: no")]
        for blocco, x, w, fondo, bordo, simbolo, alt in lati:
            self.rettangolo(sl, x, 2.22, w, 4.7, fondo=fondo, bordo=bordo, spessore=1.75, tratteggio=True,
                            raggio=0.12, nome="Riquadro")
            self.linguetta(sl, x + 0.25, 1.62, blocco["titolo"], fondo, bordo=bordo, punti=32,
                           icona=self.r.pittogramma(simbolo))
            self.blocco_testo(sl, blocco["punti"], x + 0.25, 2.45, w - 0.45, 4.3, massimo=27)
        self.note(sl, s["note"])

    def numero(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        cerchio = self.rettangolo(sl, 0.95, 1.95, 3.45, 3.45, fondo=GIALLO, forma=MSO_SHAPE.OVAL, decorativo=False,
                                  nome="Numero grande")
        tf = self.forma_con_testo(cerchio, margini=(0.0, 0.0))
        par = tf.paragraphs[0]
        par.alignment = PP_ALIGN.CENTER
        corpo = 132 if len(s["numero"]) > 1 else 190
        imposta_run(par.add_run(), s["numero"], corpo, INK, carattere=SCRITTURA)
        if s.get("didascalia"):
            self.foglietto(sl, 1.2, 5.6, 2.95, 0.8, AZZURRO, -2, s["didascalia"], punti=32, nome="Didascalia")
        soccorsi = s.get("soccorsi") or []
        h_testo = 3.05 if soccorsi else 4.9
        self.blocco_testo(sl, s["punti"], 4.85, 1.85, 7.8, h_testo, ancoraggio=MSO_ANCHOR.MIDDLE)
        for i, v in enumerate(soccorsi):
            x = 4.95 + i * 2.6
            self.cartoncino(sl, v["file"], v["alt"], x + 0.62, 4.95, 1.15, rotazione=[-2, 2, -1][i % 3])
            et = self.blocco_testo(sl, [v["testo"]], x - 0.1, 6.17, 2.6, 1.15, massimo=24, elenco=False,
                                   grassetto=True, colore=INK, nome="Soccorso")
            et.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        self.note(sl, s["note"])

    def errore(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        self.rettangolo(sl, MARGINE, 1.8, 12.0, 4.95, fondo=ROSA, bordo=ROSSO, spessore=1.75, tratteggio=True,
                        raggio=0.12, nome="Riquadro")
        self.immagine(sl, self.r.icona("avviso", ROSSO), 11.75, 1.95, 0.75, 0.75, decorativo=True, nome="Icona")
        righe = []
        for e in s["errori"]:
            righe.append(e["idea"])
            righe.append(e["correzione"])
        tb = self.casella(sl, MARGINE + 0.3, 1.95, 10.7, 4.65, nome="Errori", ancoraggio=MSO_ANCHOR.MIDDLE)
        punti = self.corpo_adatto(righe, 10.5, 4.5, 28, True, 10)
        tf = tb.text_frame
        tf.clear()
        for i, e in enumerate(s["errori"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(4)
            imposta_run(p.add_run(), tipografia(e["idea"]), punti, ROSSO, grassetto=True)
            q = tf.add_paragraph()
            q.space_after = Pt(18) if i < len(s["errori"]) - 1 else Pt(0)
            for testo, ev in pezzi(e["correzione"]):
                imposta_run(q.add_run(), testo, punti, NERO, evidenza=ev)
        self.note(sl, s["note"])

    def attivita(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s["titolo"])
        self.rettangolo(sl, MARGINE, 2.12, 8.35, 3.45, bordo=INK, tratteggio=True, raggio=0.12, nome="Riquadro")
        self.linguetta(sl, MARGINE + 0.25, 1.62, "Come si fa", AZZURRO, punti=30)
        passi = s["passi"]
        alto = 3.05 / len(passi)
        for i, riga in enumerate(passi):
            y = 2.32 + i * alto
            self.numero_tondo(sl, MARGINE + 0.25, y + (alto - 0.66) / 2, 0.66, i + 1, punti=30)
            self.blocco_testo(sl, [riga], MARGINE + 1.1, y, 7.1, alto, massimo=26, elenco=False,
                              ancoraggio=MSO_ANCHOR.MIDDLE, nome=f"Passo {i + 1}")
        rif = s["riferimento"]
        url = self.dati["sito"] + rif["url"]
        tb = self.casella(sl, MARGINE, 5.7, 8.5, 1.5, nome="Materiale")
        tf = tb.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.space_after = Pt(6)
        imposta_run(p.add_run(), "Scheda: ", 24, NERO, grassetto=True)
        imposta_run(p.add_run(), tipografia(rif["testo"]), 24, INK, link=url)
        q = tf.add_paragraph()
        imposta_run(q.add_run(), tipografia(s["materiale"]), 24, NERO)
        altezza = self.altezza_paragrafi(["Scheda: " + rif["testo"], s["materiale"]], 8.34, 24, spazio_dopo=6)
        if altezza > 1.4:
            raise ErroreSlide(f"Lezione {self.lez['numero']}, slide {self.indice}: materiale troppo lungo")
        self.foglietto(sl, 9.5, 1.85, 3.05, 1.2, GIALLO, 3, s["durata"], punti=40, nome="Durata")
        im = s["immagine"]
        self.cartoncino(sl, im["file"], im["alt"], 9.8, 3.55, 2.7, rotazione=-2)
        self.note(sl, s["note"], extra=[f"Scheda da distribuire: {url}"])

    def ricordiamo(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s.get("titolo", "Ricordiamo"))
        self.rettangolo(sl, MARGINE, 1.75, 12.0, 5.0, fondo=VERDE_CHIARO, bordo=VERDE, spessore=1.75,
                        tratteggio=True, raggio=0.12, nome="Riquadro")
        punti = s["punti"]
        alto = 4.6 / len(punti)
        for i, v in enumerate(punti):
            y = 1.95 + i * alto
            lato = min(1.3, alto - 0.2)
            self.cartoncino(sl, v["file"], v["alt"], MARGINE + 0.35, y + (alto - lato) / 2, lato,
                            rotazione=[-2, 2, -1][i % 3])
            self.blocco_testo(sl, [v["testo"]], MARGINE + 2.0, y, 9.7, alto, massimo=30, elenco=False,
                              grassetto=True, ancoraggio=MSO_ANCHOR.MIDDLE, nome=f"Ricordiamo {i + 1}")
        self.note(sl, s["note"])

    def fonti(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s.get("titolo", "Per saperne di più"))
        self.rettangolo(sl, MARGINE, 2.15, 12.0, 4.95, bordo=INK, tratteggio=True, raggio=0.12, nome="Riquadro")
        self.linguetta(sl, MARGINE + 0.25, 1.65, "Le pagine del sito", AZZURRO, punti=30)
        tb = self.casella(sl, MARGINE + 0.3, 2.35, 11.4, 4.6, nome="Fonti")
        tf = tb.text_frame
        tf.clear()
        voci = list(self.lez["fonti"])
        perc = self.dati["percorso"]
        voci.append({"testo": f"Percorsi didattici: Percorso {perc['numero']}", "url": perc["pagina"]})
        sito = self.dati["sito"]
        for i, f in enumerate(voci):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(4)
            puntato(p)
            imposta_run(p.add_run(), tipografia(f["testo"]), 24, INK, link=sito + f["url"])
        p = tf.add_paragraph()
        p.space_before = Pt(10)
        imposta_run(p.add_run(), "Tutte le pagine sono su protezionecivilegenzano.it.", 24, NERO)
        altezza = self.altezza_paragrafi([f["testo"] for f in voci] + ["x"], 11.24 - 0.42, 24, spazio_dopo=4)
        if altezza > 4.6:
            raise ErroreSlide(f"Lezione {self.lez['numero']}: troppe fonti per una slide")
        extra = ["Indirizzi completi:"] + [f"- {f['testo']}: {sito}{f['url']}" for f in voci]
        self.note(sl, s["note"], extra=extra)

    def chiusura(self, s):
        sl = self.nuova_slide()
        self.titolo(sl, s.get("titolo", "Grazie!"), punti=48)
        self.blocco_testo(sl, self.dati["crediti"], MARGINE, 1.8, 12.0, 2.8, massimo=24, elenco=False,
                          spazio_dopo=6, nome="Crediti")
        (esc, alt_esc), (fep, alt_fep), (snpc, alt_snpc) = AFFILIAZIONI
        self.immagine(sl, esc, 0.85, 4.85, 1.8, 1.8, alt=alt_esc, nome="Logo Quality Label")
        tb = self.casella(sl, 2.8, 4.95, 4.1, 1.6, nome="Codice di accreditamento", ancoraggio=MSO_ANCHOR.MIDDLE)
        tf = tb.text_frame
        tf.clear()
        imposta_run(tf.paragraphs[0].add_run(), "Corpo europeo di solidarietà", 24, NERO)
        q = tf.add_paragraph()
        imposta_run(q.add_run(), f"Codice {CODICE_ESC}", 26, INK, grassetto=True)
        self.immagine(sl, fep, 7.2, 4.85, 1.8, 1.8, alt=alt_fep, nome="Logo FE.PI.VOL.")
        self.immagine(sl, snpc, 9.8, 4.75, 1.65, 2.0, alt=alt_snpc, nome="Logo SNPC Volontariato")
        self.note(sl, s["note"])

    # ── montaggio ──
    def costruisci(self) -> bytes:
        for s in self.lez["slide"]:
            metodo = getattr(self, s["tipo"], None)
            if metodo is None:
                raise ErroreSlide(f"Tipo di slide sconosciuto: {s['tipo']}")
            metodo(s)
        self._proprieta()
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()

    def _proprieta(self):
        cp = self.prs.core_properties
        perc = self.dati["percorso"]
        quando = dt.datetime.combine(self.dati["aggiornato"], dt.time(9, 0))
        cp.title = tipografia(f"Lezione {self.lez['numero']}: {self.lez['titolo']}")
        cp.subject = f"Percorso {perc['numero']} · {perc['titolo']} · {perc['fascia']}"
        cp.author = self.dati["autore"]
        cp.last_modified_by = self.dati["autore"]
        cp.comments = "Slide per chi conduce le lezioni nelle scuole, con le note per chi parla in ogni slide."
        cp.keywords = "protezione civile; scuola primaria; educazione civica; Genzano di Roma"
        cp.category = "Formazione"
        cp.language = "it-IT"
        cp.created = quando
        cp.modified = quando
        cp.revision = 1
        cp.last_printed = quando


# ───────────────────────────── pacchetto finale ─────────────────────────────

def eot(ttf: Path) -> bytes:
    """Carattere TrueType nel formato Embedded OpenType (versione 0x00020001, non compresso)."""
    dati = ttf.read_bytes()
    font = TTFont(ttf)
    os2 = font["OS/2"]

    def nome(i):
        for rec in font["name"].names:
            if rec.nameID == i and rec.platformID == 3:
                return rec.toUnicode()
        return ""

    def stringa(s):
        b = s.encode("utf-16-le")
        return struct.pack("<H", len(b)) + b

    pan = os2.panose
    corpo = bytes([pan.bFamilyType, pan.bSerifStyle, pan.bWeight, pan.bProportion, pan.bContrast,
                   pan.bStrokeVariation, pan.bArmStyle, pan.bLetterForm, pan.bMidline, pan.bXHeight])
    corpo += struct.pack("<BB", 1, 1 if os2.fsSelection & 1 else 0)
    corpo += struct.pack("<IHH", os2.usWeightClass, os2.fsType, 0x504C)
    corpo += struct.pack("<IIII", os2.ulUnicodeRange1, os2.ulUnicodeRange2, os2.ulUnicodeRange3, os2.ulUnicodeRange4)
    corpo += struct.pack("<II", os2.ulCodePageRange1, os2.ulCodePageRange2)
    corpo += struct.pack("<I", font["head"].checkSumAdjustment)
    corpo += struct.pack("<IIII", 0, 0, 0, 0)
    for i in (1, 2, 5, 4):
        corpo += struct.pack("<H", 0) + stringa(nome(i))
    corpo += struct.pack("<HH", 0, 0)  # Padding5, RootStringSize
    intestazione = 16 + len(corpo)
    return struct.pack("<IIII", intestazione + len(dati), len(dati), 0x00020001, 0) + corpo + dati


APP_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<TotalTime>0</TotalTime><PresentationFormat>Widescreen</PresentationFormat>'
           '<Slides>{n}</Slides><Notes>{n}</Notes><HiddenSlides>0</HiddenSlides><MMClips>0</MMClips>'
           '<ScaleCrop>false</ScaleCrop><Company>{societa}</Company><LinksUpToDate>false</LinksUpToDate>'
           '<SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion>'
           '</Properties>')


def confeziona(pptx: bytes, ttf: Path, n_slide: int, autore: str, data: dt.date) -> bytes:
    """Incorpora il carattere a mano, riscrive le proprietà dell'applicazione e fissa le date nel pacchetto."""
    zin = zipfile.ZipFile(io.BytesIO(pptx))
    parti = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    parti["ppt/fonts/font1.fntdata"] = eot(ttf)
    ct = parti["[Content_Types].xml"].decode("utf-8")
    if 'Extension="fntdata"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="fntdata" ContentType="application/x-fontdata"/><Default ', 1)
    parti["[Content_Types].xml"] = ct.encode("utf-8")
    rels = parti["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    rels = rels.replace("</Relationships>",
                        '<Relationship Id="rIdCarattere1" Type="http://schemas.openxmlformats.org/officeDocument/'
                        '2006/relationships/font" Target="fonts/font1.fntdata"/></Relationships>')
    parti["ppt/_rels/presentation.xml.rels"] = rels.encode("utf-8")
    pres = parti["ppt/presentation.xml"].decode("utf-8")
    lista = ('<p:embeddedFontLst><p:embeddedFont><p:font typeface="%s" pitchFamily="66" charset="0"/>'
             '<p:regular r:id="rIdCarattere1"/></p:embeddedFont></p:embeddedFontLst>' % SCRITTURA)
    pres, n = re.subn(r"(<p:notesSz[^>]*/>)", r"\1" + lista, pres, count=1)
    if n != 1:
        raise ErroreSlide("Struttura di presentation.xml inattesa: carattere non incorporato")
    pres = pres.replace("<p:presentation ", '<p:presentation embedTrueTypeFonts="1" ', 1)
    parti["ppt/presentation.xml"] = pres.encode("utf-8")
    parti["docProps/app.xml"] = APP_XML.format(n=n_slide, societa=autore).encode("utf-8")
    quando = (data.year, data.month, data.day, 9, 0, 0)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        ordine = ["[Content_Types].xml"] + [n for n in parti if n != "[Content_Types].xml"]
        for nome in ordine:
            info = zipfile.ZipInfo(nome, date_time=quando)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, parti[nome])
    return out.getvalue()


def nome_file(dati, lez, estensione):
    return f"percorso-{dati['percorso']['numero']}-lezione-{lez['numero']}-{lez['slug']}.{estensione}"


# ───────────────────────────── PDF ─────────────────────────────

def converti_pdf(pptx: Path, ttf: Path, dati: dict, lez: dict) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("LibreOffice non trovato: il PDF non si può generare.")
    with tempfile.TemporaryDirectory(prefix="slide-pdf-") as tmp:
        tmp = Path(tmp)
        caratteri = tmp / "caratteri"
        caratteri.mkdir()
        shutil.copy(ttf, caratteri)
        conf = tmp / "fonts.conf"
        conf.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
                        '<include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                        f'<dir>{caratteri}</dir><cachedir>{tmp / "cache"}</cachedir></fontconfig>', encoding="utf-8")
        env = os.environ.copy()
        env["FONTCONFIG_FILE"] = str(conf)
        env["SAL_USE_VCLPLUGIN"] = "svp"
        filtro = 'pdf:impress_pdf_Export:{"UseTaggedPDF":{"type":"boolean","value":"true"},' \
                 '"ExportNotesPages":{"type":"boolean","value":"false"}}'
        subprocess.run([soffice, f"-env:UserInstallation={(tmp / 'profilo').as_uri()}", "--headless",
                        "--convert-to", filtro, "--outdir", str(tmp), str(pptx)],
                       env=env, check=True, capture_output=True, timeout=600)
        grezzo = tmp / (pptx.stem + ".pdf")
        if not grezzo.exists():
            raise SystemExit(f"Conversione in PDF non riuscita: {pptx.name}")
        finale = pptx.with_suffix(".pdf")
        try:
            import pymupdf
        except ImportError:
            shutil.copy(grezzo, finale)
            print("  Avviso: pymupdf assente, metadati del PDF non ripuliti.")
            return finale
        doc = pymupdf.open(grezzo)
        perc = dati["percorso"]
        doc.set_metadata({
            "title": tipografia(f"Lezione {lez['numero']}: {lez['titolo']}"),
            "author": dati["autore"],
            "subject": f"Percorso {perc['numero']} · {perc['titolo']} · {perc['fascia']}",
            "keywords": "protezione civile; scuola primaria; educazione civica; Genzano di Roma",
            "creator": dati["autore"],
            "producer": "",
        })
        doc.set_language("it-IT")
        doc.save(finale, garbage=3, deflate=True)
        doc.close()
        return finale


# ───────────────────────────── pagine del sito ─────────────────────────────

def etichetta_mb(percorso: Path) -> str:
    mb = percorso.stat().st_size / 1048576
    return ("%.1f" % max(mb, 0.1)).replace(".", ",") + " MB"


def aggiorna_pagine():
    for pagina in PAGINE_CON_LINK:
        if not pagina.exists():
            continue
        testo = pagina.read_text(encoding="utf-8")

        def sostituisci(m):
            etichetta, rel = m.group(1), m.group(2)
            f = ROOT / "static" / rel.lstrip("/")
            if not f.exists():
                return m.group(0)
            return f"[{etichetta} ({etichetta_mb(f)})]({rel})"

        nuovo = re.sub(r"\[(PowerPoint|PDF) \([^)]*\)\]\((/formazione/slide-lezioni/[^)\s]+)\)", sostituisci, testo)
        if nuovo != testo:
            pagina.write_text(nuovo, encoding="utf-8")
            print(f"  Dimensioni aggiornate in {pagina.relative_to(ROOT)}")


# ───────────────────────────── verifica ─────────────────────────────

def _impronta(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def verifica(dati: dict, solo: int | None) -> int:
    errori: list[str] = []
    misura_t = Misura(MISURA_ARIAL)
    misura_b = Misura(MISURA_ARIAL_B)
    loghi = {_impronta(LOGO_GRUPPO): "gruppo"}
    for p, _ in AFFILIAZIONI:
        loghi[_impronta(p)] = p.stem
    for lez in dati["lezioni"]:
        if solo and lez["numero"] != solo:
            continue
        f = USCITA / nome_file(dati, lez, "pptx")
        tag = f.name
        if not f.exists():
            errori.append(f"{tag}: file mancante")
            continue
        grezzo = f.read_bytes()
        with zipfile.ZipFile(io.BytesIO(grezzo)) as z:
            for nome in z.namelist():
                if nome.endswith((".xml", ".rels")):
                    contenuto = z.read(nome).decode("utf-8", "replace").lower()
                    for parola in PAROLE_VIETATE:
                        if parola in contenuto:
                            errori.append(f"{tag}: «{parola}» in {nome}")
            if not any(n.startswith("ppt/fonts/") for n in z.namelist()):
                errori.append(f"{tag}: carattere a mano non incorporato")
        prs = Presentation(io.BytesIO(grezzo))
        if prs.core_properties.author != dati["autore"]:
            errori.append(f"{tag}: autore del documento diverso da «{dati['autore']}»")
        slides = list(prs.slides)
        if len(slides) != len(lez["slide"]):
            errori.append(f"{tag}: {len(slides)} slide invece di {len(lez['slide'])}")
        if not 10 <= len(slides) <= 14:
            errori.append(f"{tag}: {len(slides)} slide, ne servono da 10 a 14")
        minuti = 0
        posti_loghi: list[tuple[int, str]] = []
        for i, sl in enumerate(slides, 1):
            dove = f"{tag}, slide {i}"
            titolo = sl.shapes.title
            if titolo is None or not titolo.text_frame.text.strip():
                errori.append(f"{dove}: manca il titolo nel segnaposto del titolo")
            for sh in sl.shapes:
                if sh.shape_type == 13:  # immagine
                    el = cnvpr(sh)
                    decorativo = el.find(qn("a:extLst")) is not None and DECORATIVO_URI in etree.tostring(el).decode()
                    if not el.get("descr") and not decorativo:
                        errori.append(f"{dove}: immagine «{sh.name}» senza testo alternativo")
                    imp = hashlib.sha256(sh.image.blob).hexdigest()
                    if imp in loghi:
                        posti_loghi.append((i, loghi[imp]))
                if not sh.has_text_frame:
                    continue
                e_titolo = titolo is not None and sh.shape_id == titolo.shape_id
                minimo = MIN_TITOLO if e_titolo else MIN_TESTO
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if not r.text.strip():
                            continue
                        if r.font.size is None or r.font.size.pt < minimo:
                            corpo = r.font.size.pt if r.font.size else "non indicato"
                            errori.append(f"{dove}: «{r.text[:30]}» a {corpo} punti (minimo {minimo})")
                if not e_titolo and sh.has_text_frame and sh.rotation == 0:
                    errori += _controlla_spazio(sh, dove, misura_t, misura_b)
            note = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
            m = re.search(r"Tempo: (\d+) minut", note)
            if not note.strip() or "Cosa dire:" not in note or not m:
                errori.append(f"{dove}: note per chi parla incomplete")
            else:
                minuti += int(m.group(1))
        if minuti != 60:
            errori.append(f"{tag}: i tempi delle note fanno {minuti} minuti invece di 60")
        ultimo = len(slides)
        attesi = sorted([(1, "gruppo")] + [(ultimo, p.stem) for p, _ in AFFILIAZIONI])
        if sorted(posti_loghi) != attesi:
            errori.append(f"{tag}: loghi fuori posto o in più: {sorted(posti_loghi)}")
        if CODICE_ESC not in "".join(sh.text_frame.text for sh in slides[-1].shapes if sh.has_text_frame):
            errori.append(f"{tag}: manca il codice {CODICE_ESC} accanto al Quality Label")
        pdf = f.with_suffix(".pdf")
        if pdf.exists():
            try:
                import pymupdf
                doc = pymupdf.open(pdf)
                if doc.page_count != len(slides):
                    errori.append(f"{pdf.name}: {doc.page_count} pagine invece di {len(slides)}")
                caratteri = {c[3] for pg in doc for c in pg.get_fonts()}
                if not any("EduNSWACTCursive" in c for c in caratteri):
                    errori.append(f"{pdf.name}: il carattere a mano non è incorporato")
                meta = doc.metadata or {}
                if meta.get("author") != dati["autore"]:
                    errori.append(f"{pdf.name}: autore del PDF diverso da «{dati['autore']}»")
                testo_meta = json.dumps(meta).lower()
                for parola in PAROLE_VIETATE:
                    if parola in testo_meta:
                        errori.append(f"{pdf.name}: «{parola}» nei metadati")
            except ImportError:
                pass
    for e in errori:
        print(f"ERRORE  {e}")
    if not errori:
        print("Slide verificate: titoli, note, tempi, testi alternativi, corpi, spazi, loghi e metadati a posto.")
    return len(errori)


def _controlla_spazio(sh, dove, misura_t, misura_b) -> list[str]:
    """Stima se il testo entra nella sua casella (stesse misure usate per costruirla)."""
    tf = sh.text_frame
    larg = Emu(sh.width - tf.margin_left - tf.margin_right).inches
    alto = Emu(sh.height - tf.margin_top - tf.margin_bottom).inches
    tot = 0.0
    paragrafi = [p for p in tf.paragraphs if p.runs]
    for k, p in enumerate(paragrafi):
        corpo = max(r.font.size.pt for r in p.runs if r.font.size)
        grassetto = all(r.font.bold for r in p.runs)
        scrittura = any(r.font.name == SCRITTURA for r in p.runs)
        if scrittura:
            continue
        ppr = p._p.pPr
        rientro = Emu(int(ppr.get("marL"))).inches if ppr is not None and ppr.get("marL") else 0.0
        testo = "".join(r.text for r in p.runs)
        righe = righe_a_capo(testo, larg - rientro, corpo, misura_b if grassetto else misura_t)
        tot += len(righe) * corpo * Lezione.LH_TESTO / 72
        if k < len(paragrafi) - 1 and p.space_after is not None:
            tot += p.space_after.pt / 72
    if tot > alto + 0.02:
        return [f"{dove}: il testo di «{sh.name}» non entra ({tot:.2f} pollici su {alto:.2f})"]
    return []


# ───────────────────────────── main ─────────────────────────────

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pdf", action="store_true", help="converte anche in PDF con LibreOffice")
    ap.add_argument("--verifica", action="store_true", help="controlla i file già generati, senza scrivere")
    ap.add_argument("--solo", type=int, help="numero della lezione da trattare")
    args = ap.parse_args()
    dati = yaml.safe_load(DATI.read_text(encoding="utf-8"))
    if args.verifica:
        return verifica(dati, args.solo)
    USCITA.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="slide-lezioni-") as tmp:
        risorse = Risorse(Path(tmp))
        for lez in dati["lezioni"]:
            if args.solo and lez["numero"] != args.solo:
                continue
            try:
                grezzo = Lezione(dati, lez, risorse).costruisci()
            except ErroreSlide as e:
                print(f"ERRORE  {e}")
                return 1
            finale = confeziona(grezzo, risorse.ttf, len(lez["slide"]), dati["autore"], dati["aggiornato"])
            out = USCITA / nome_file(dati, lez, "pptx")
            out.write_bytes(finale)
            print(f"  {out.relative_to(ROOT)} ({etichetta_mb(out)}, {len(lez['slide'])} slide)")
            if args.pdf:
                pdf = converti_pdf(out, risorse.ttf, dati, lez)
                print(f"  {pdf.relative_to(ROOT)} ({etichetta_mb(pdf)})")
    aggiorna_pagine()
    return 0


if __name__ == "__main__":
    sys.exit(main())
